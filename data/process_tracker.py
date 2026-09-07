"""
TRACES Automation — Process Tracker (Excel)

Manages:
    Process_Files/automation_tracker.xlsx
    Process_Files/failed_records.xlsx

These are internal operational files, separate from the business Master Tracker.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.worksheet import Worksheet

from config import AUTOMATION_TRACKER_FILE, FAILED_RECORDS_FILE
from utils.logger import logger

# ──────────────────────────────────────────────
# Column definitions
# ──────────────────────────────────────────────
AUTOMATION_COLUMNS = [
    "SL_No",
    "Unique_Key",
    "Certificate_Type",
    "Financial_Year",
    "PAN",
    "Certificate_Number",
    "Status",
    "ARN_Request_Number",
    "Initiated_At",
    "Downloaded_At",
    "Original_Filename",
    "Final_Filename",
    "PDF_Path",
    "Extraction_Status",
    "Master_Update_Status",
    "Services_Update_Status",
    "Retry_Count",
    "Error_Message",
    "Last_Updated",
]

FAILED_COLUMNS = [
    "SL_No",
    "Unique_Key",
    "Certificate_Type",
    "Financial_Year",
    "PAN",
    "Certificate_Number",
    "Status",
    "Error_Message",
    "Retry_Count",
    "Last_Attempt_At",
    "Created_At",
]


def _create_workbook(path: Path, columns: List[str]) -> None:
    """Create an Excel workbook with styled headers."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws: Worksheet = wb.active  # type: ignore
    ws.title = "Tracker"

    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) + 2, 12)

    ws.freeze_panes = "A2"
    wb.save(str(path))
    logger.debug("Created workbook: %s", path)


class ProcessTracker:
    """Manages the automation_tracker.xlsx workbook."""

    def __init__(self, path: Optional[Path] = None) -> None:
        self.path = path or AUTOMATION_TRACKER_FILE
        if not self.path.exists():
            _create_workbook(self.path, AUTOMATION_COLUMNS)

    def append_record(self, data: Dict[str, str]) -> None:
        """Append a single row from a dict."""
        wb = load_workbook(str(self.path))
        ws = wb["Tracker"]
        new_row = ws.max_row + 1
        data["SL_No"] = str(new_row - 1)
        data["Last_Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for col_idx, col_name in enumerate(AUTOMATION_COLUMNS, start=1):
            ws.cell(row=new_row, column=col_idx, value=data.get(col_name, ""))
        wb.save(str(self.path))

    def update_record(self, unique_key: str, updates: Dict[str, str]) -> bool:
        """Find row by unique_key and update columns. Returns True if found."""
        wb = load_workbook(str(self.path))
        ws = wb["Tracker"]
        key_col = AUTOMATION_COLUMNS.index("Unique_Key") + 1
        updates["Last_Updated"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        for row_num in range(2, ws.max_row + 1):
            if str(ws.cell(row=row_num, column=key_col).value or "") == unique_key:
                for col_name, value in updates.items():
                    if col_name in AUTOMATION_COLUMNS:
                        col_idx = AUTOMATION_COLUMNS.index(col_name) + 1
                        ws.cell(row=row_num, column=col_idx, value=value)
                wb.save(str(self.path))
                return True
        # Not found — append instead
        updates["Unique_Key"] = unique_key
        self.append_record(updates)
        return False


class FailedRecordsTracker:
    """Manages the failed_records.xlsx workbook."""

    def __init__(self, path: Optional[Path] = None) -> None:
        self.path = path or FAILED_RECORDS_FILE
        if not self.path.exists():
            _create_workbook(self.path, FAILED_COLUMNS)

    def add_failure(self, data: Dict[str, str]) -> None:
        """Append a failure record."""
        wb = load_workbook(str(self.path))
        ws = wb["Tracker"]
        new_row = ws.max_row + 1
        data["SL_No"] = str(new_row - 1)
        data["Created_At"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        for col_idx, col_name in enumerate(FAILED_COLUMNS, start=1):
            ws.cell(row=new_row, column=col_idx, value=data.get(col_name, ""))
        wb.save(str(self.path))
        logger.debug("Recorded failure: %s", data.get("Unique_Key", ""))
