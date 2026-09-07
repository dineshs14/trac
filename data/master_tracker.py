"""
TRACES Automation — Master Tracker (Excel)

Manages the single business workbook:
    Master_Tracker/Lower_Nil_TDS_Master_Tracker.xlsx

Sheet: Certificate_Data

All Excel operations are centralised here.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from config import MASTER_TRACKER_FILE, MASTER_COLUMNS, MASTER_SHEET_NAME
from data.models import CertificateRecord
from utils.logger import logger


class MasterTracker:
    """Centralised read/write manager for the Master Excel workbook."""

    def __init__(self, path: Optional[Path] = None) -> None:
        self.path = path or MASTER_TRACKER_FILE
        self._ensure_workbook()

    # ── Initialisation ──────────────────────────

    def _ensure_workbook(self) -> None:
        """Create the workbook and header row if the file does not exist."""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            logger.debug("Master tracker exists: %s", self.path)
            return

        wb = Workbook()
        ws: Worksheet = wb.active  # type: ignore[assignment]
        ws.title = MASTER_SHEET_NAME

        # Header row
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")

        for col_idx, col_name in enumerate(MASTER_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        # Freeze header
        ws.freeze_panes = "A2"

        # Auto-width approximation
        for col_idx, col_name in enumerate(MASTER_COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
                len(col_name) + 4, 14
            )

        wb.save(str(self.path))
        logger.info("Created Master Tracker: %s", self.path)

    # ── Key helpers ─────────────────────────────

    @staticmethod
    def _match_key(record: Dict[str, Any]) -> str:
        """Build the composite key for row matching."""
        return (
            f"{record.get('Certificate_Type', '')}|"
            f"{record.get('Financial_Year', '')}|"
            f"{record.get('PAN', '')}|"
            f"{record.get('Certificate_Number', '')}"
        )

    def _col_index(self, col_name: str) -> int:
        """Return the 1-based column index for a column name."""
        return MASTER_COLUMNS.index(col_name) + 1

    # ── Read operations ─────────────────────────

    def _load_ws(self) -> tuple:
        """Load and return (workbook, worksheet)."""
        wb = load_workbook(str(self.path))
        ws = wb[MASTER_SHEET_NAME]
        return wb, ws

    def find_row(
        self, cert_type: str, fy: str, pan: str, cert_no: str
    ) -> Optional[int]:
        """Find the row number (1-based) matching the composite key.

        Returns None if no match is found.
        """
        _, ws = self._load_ws()
        type_col = self._col_index("Certificate_Type")
        fy_col = self._col_index("Financial_Year")
        pan_col = self._col_index("PAN")
        cert_col = self._col_index("Certificate_Number")

        for row_num in range(2, ws.max_row + 1):
            if (
                str(ws.cell(row=row_num, column=type_col).value or "") == cert_type
                and str(ws.cell(row=row_num, column=fy_col).value or "") == fy
                and str(ws.cell(row=row_num, column=pan_col).value or "") == pan
                and str(ws.cell(row=row_num, column=cert_col).value or "") == cert_no
            ):
                return row_num
        return None

    def find_row_by_pan_cert_fy(
        self, pan: str, cert_no: str, fy: str
    ) -> Optional[int]:
        """Find row matching PAN + Certificate Number + FY (ignoring cert type).

        Used by the Services updater where certificate type might not be directly available.
        """
        _, ws = self._load_ws()
        fy_col = self._col_index("Financial_Year")
        pan_col = self._col_index("PAN")
        cert_col = self._col_index("Certificate_Number")

        for row_num in range(2, ws.max_row + 1):
            if (
                str(ws.cell(row=row_num, column=fy_col).value or "") == fy
                and str(ws.cell(row=row_num, column=pan_col).value or "") == pan
                and str(ws.cell(row=row_num, column=cert_col).value or "") == cert_no
            ):
                return row_num
        return None

    def get_all_pans(self, fy: str) -> List[str]:
        """Return distinct PANs for a given FY."""
        _, ws = self._load_ws()
        fy_col = self._col_index("Financial_Year")
        pan_col = self._col_index("PAN")
        pans = set()
        for row_num in range(2, ws.max_row + 1):
            if str(ws.cell(row=row_num, column=fy_col).value or "") == fy:
                pan_val = str(ws.cell(row=row_num, column=pan_col).value or "")
                if pan_val:
                    pans.add(pan_val)
        return sorted(pans)

    def count_for_pan(self, pan: str, fy: str) -> int:
        """Count certificates for a specific PAN in a given FY."""
        _, ws = self._load_ws()
        fy_col = self._col_index("Financial_Year")
        pan_col = self._col_index("PAN")
        count = 0
        for row_num in range(2, ws.max_row + 1):
            if (
                str(ws.cell(row=row_num, column=fy_col).value or "") == fy
                and str(ws.cell(row=row_num, column=pan_col).value or "") == pan
            ):
                count += 1
        return count

    # ── Write operations ────────────────────────

    def _next_sl_no(self, ws: Worksheet) -> int:
        """Determine the next serial number."""
        sl_col = self._col_index("SL_No")
        max_sl = 0
        for row_num in range(2, ws.max_row + 1):
            val = ws.cell(row=row_num, column=sl_col).value
            if val and isinstance(val, (int, float)):
                max_sl = max(max_sl, int(val))
        return max_sl + 1

    def upsert_record(self, record: CertificateRecord) -> str:
        """Insert or update a certificate record in the Master Tracker.

        Returns 'ADDED' or 'UPDATED'.
        """
        wb, ws = self._load_ws()
        existing_row = self.find_row(
            record.certificate_type,
            record.financial_year,
            record.pan,
            record.certificate_number,
        )

        data = self._record_to_dict(record)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        data["Last_Updated"] = now

        if existing_row:
            # Update existing row — do not overwrite SL_No
            for col_name in MASTER_COLUMNS:
                if col_name == "SL_No":
                    continue
                col_idx = self._col_index(col_name)
                new_val = data.get(col_name)
                if new_val is not None and new_val != "":
                    ws.cell(row=existing_row, column=col_idx, value=new_val)
            wb.save(str(self.path))
            logger.debug("Master updated row %d: %s", existing_row, record.unique_key)
            return "UPDATED"
        else:
            # Append new row
            new_row = ws.max_row + 1
            data["SL_No"] = self._next_sl_no(ws)
            # Derive count
            data["Count"] = self.count_for_pan(record.pan, record.financial_year) + 1
            data["Date_Received"] = now
            for col_name in MASTER_COLUMNS:
                col_idx = self._col_index(col_name)
                ws.cell(row=new_row, column=col_idx, value=data.get(col_name, ""))
            wb.save(str(self.path))
            logger.debug("Master added row %d: %s", new_row, record.unique_key)
            return "ADDED"

    def update_services_data(
        self,
        pan: str,
        cert_no: str,
        fy: str,
        updates: Dict[str, Any],
    ) -> bool:
        """Update specific columns for an existing row matched by PAN+Cert+FY.

        Used by the Services updater module.
        Returns True if a matching row was found and updated.
        """
        wb, ws = self._load_ws()
        row_num = self.find_row_by_pan_cert_fy(pan, cert_no, fy)
        if row_num is None:
            logger.warning(
                "Services update: no matching row for PAN=%s, Cert=%s, FY=%s",
                pan, cert_no, fy,
            )
            return False

        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        updates["Last_Updated"] = now

        for col_name, value in updates.items():
            if col_name in MASTER_COLUMNS:
                col_idx = self._col_index(col_name)
                ws.cell(row=row_num, column=col_idx, value=value)

        wb.save(str(self.path))
        logger.debug(
            "Services data updated row %d: PAN=%s, Cert=%s",
            row_num, pan, cert_no,
        )
        return True

    # ── Conversion helper ───────────────────────

    @staticmethod
    def _record_to_dict(rec: CertificateRecord) -> Dict[str, Any]:
        """Convert a CertificateRecord to a dict keyed by column names."""
        return {
            "Certificate_Type": rec.certificate_type,
            "PAN": rec.pan,
            "Vendor_Name": rec.vendor_name,
            "Certificate_Number": rec.certificate_number,
            "Certificate_Short_Number": rec.certificate_short_number,
            "Financial_Year": rec.financial_year,
            "TDS_Rate": rec.tds_rate,
            "Nature_of_Payment": rec.nature_of_payment,
            "Section": rec.section,
            "Section_Code": rec.section_code,
            "Certificate_Limit": rec.certificate_limit,
            "Valid_From": rec.valid_from,
            "Valid_To": rec.valid_to,
            "Certificate_Name": rec.certificate_name,
            "Date_of_Issue": rec.date_of_issue,
            "Application_Form_No": rec.application_form_no,
            "Applicable_Income_Tax_Act": rec.applicable_income_tax_act,
            "Q1_Amount_Consumed": rec.q1_amount_consumed,
            "Q2_Amount_Consumed": rec.q2_amount_consumed,
            "Q3_Amount_Consumed": rec.q3_amount_consumed,
            "Q4_Amount_Consumed": rec.q4_amount_consumed,
            "Total_Amount_Consumed": rec.total_amount_consumed,
            "Available_Amount": rec.available_amount,
            "Date_of_Cancellation": rec.date_of_cancellation,
            "Notes": rec.notes,
            "PDF_Path": rec.pdf_path,
            "Processing_Status": rec.processing_status,
            "Last_Updated": rec.last_updated,
        }
