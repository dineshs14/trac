"""
Tests for data/master_tracker.py — Excel Master Tracker operations.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
from data.master_tracker import MasterTracker
from data.models import CertificateRecord
from config import MASTER_COLUMNS, MASTER_SHEET_NAME


@pytest.fixture
def tracker(tmp_path):
    """Create a temporary Master Tracker workbook."""
    path = tmp_path / "test_tracker.xlsx"
    return MasterTracker(path=path)


class TestMasterTracker:
    """Tests for the MasterTracker class."""

    def test_creates_workbook(self, tracker):
        assert tracker.path.exists()

    def test_has_correct_sheet(self, tracker):
        from openpyxl import load_workbook

        wb = load_workbook(str(tracker.path))
        assert MASTER_SHEET_NAME in wb.sheetnames

    def test_has_all_columns(self, tracker):
        from openpyxl import load_workbook

        wb = load_workbook(str(tracker.path))
        ws = wb[MASTER_SHEET_NAME]
        for col_idx, col_name in enumerate(MASTER_COLUMNS, start=1):
            assert ws.cell(row=1, column=col_idx).value == col_name

    def test_upsert_new_record(self, tracker):
        record = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Test Vendor",
            certificate_number="1NA0926KBC",
            financial_year="2026-27",
            tds_rate="5%",
        )
        action = tracker.upsert_record(record)
        assert action == "ADDED"

    def test_upsert_update_existing(self, tracker):
        record = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Test Vendor",
            certificate_number="1NA0926KBC",
            financial_year="2026-27",
            tds_rate="5%",
        )
        tracker.upsert_record(record)

        # Update the same record with new data
        record.tds_rate = "10%"
        record.nature_of_payment = "Rent"
        action = tracker.upsert_record(record)
        assert action == "UPDATED"

    def test_no_duplicate_rows(self, tracker):
        record = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Test Vendor",
            certificate_number="1NA0926KBC",
            financial_year="2026-27",
        )
        tracker.upsert_record(record)
        tracker.upsert_record(record)
        tracker.upsert_record(record)

        from openpyxl import load_workbook

        wb = load_workbook(str(tracker.path))
        ws = wb[MASTER_SHEET_NAME]
        # Should have header + 1 data row
        assert ws.max_row == 2

    def test_find_row(self, tracker):
        record = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Test Vendor",
            certificate_number="1NA0926KBC",
            financial_year="2026-27",
        )
        tracker.upsert_record(record)

        row = tracker.find_row("LOWER_TDS", "2026-27", "AAHCP9855G", "1NA0926KBC")
        assert row == 2

    def test_find_row_not_found(self, tracker):
        row = tracker.find_row("LOWER_TDS", "2026-27", "XXXXX1234Y", "NONEXIST")
        assert row is None

    def test_find_row_by_pan_cert_fy(self, tracker):
        record = CertificateRecord(
            certificate_type="CHILD",
            pan="AAHCC2532P",
            vendor_name="Child Vendor",
            certificate_number="4NA0826ACO4A023",
            financial_year="2026-27",
        )
        tracker.upsert_record(record)

        row = tracker.find_row_by_pan_cert_fy("AAHCC2532P", "4NA0826ACO4A023", "2026-27")
        assert row == 2

    def test_update_services_data(self, tracker):
        record = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Test Vendor",
            certificate_number="1NA0926KBC",
            financial_year="2026-27",
        )
        tracker.upsert_record(record)

        updated = tracker.update_services_data(
            pan="AAHCP9855G",
            cert_no="1NA0926KBC",
            fy="2026-27",
            updates={
                "Q1_Amount_Consumed": 500000,
                "Q2_Amount_Consumed": 300000,
                "Total_Amount_Consumed": 800000,
                "Date_of_Cancellation": "Not Cancelled",
            },
        )
        assert updated is True

    def test_update_services_data_not_found(self, tracker):
        updated = tracker.update_services_data(
            pan="XXXXX1234Y",
            cert_no="NONEXIST",
            fy="2026-27",
            updates={"Q1_Amount_Consumed": 100},
        )
        assert updated is False

    def test_count_for_pan(self, tracker):
        for i in range(3):
            record = CertificateRecord(
                certificate_type="LOWER_TDS",
                pan="AAHCP9855G",
                vendor_name="Test Vendor",
                certificate_number=f"CERT{i}",
                financial_year="2026-27",
            )
            tracker.upsert_record(record)

        assert tracker.count_for_pan("AAHCP9855G", "2026-27") == 3
        assert tracker.count_for_pan("NONEXIST", "2026-27") == 0

    def test_mixed_cert_types(self, tracker):
        lower = CertificateRecord(
            certificate_type="LOWER_TDS",
            pan="AAHCP9855G",
            vendor_name="Lower Vendor",
            certificate_number="CERT_LOWER",
            financial_year="2026-27",
        )
        child = CertificateRecord(
            certificate_type="CHILD",
            pan="AAHCP9855G",
            vendor_name="Child Vendor",
            certificate_number="CERT_CHILD",
            financial_year="2026-27",
        )
        tracker.upsert_record(lower)
        tracker.upsert_record(child)

        from openpyxl import load_workbook

        wb = load_workbook(str(tracker.path))
        ws = wb[MASTER_SHEET_NAME]
        assert ws.max_row == 3  # header + 2 data rows
